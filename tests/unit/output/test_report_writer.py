# SPDX-FileCopyrightText: 2026 Fonden Mærsk Mc-Kinney Møller Center for Zero Carbon Shipping
# SPDX-License-Identifier: Apache-2.0

"""
Resolution of Report property requests against the node registry, and the per-sheet error
containment of write_report.
"""

from __future__ import annotations

import logging
from types import SimpleNamespace

import numpy as np
import openpyxl as xl

from navigate.core.enum_ import ReportReduceID
from navigate.core.node_registry import Nodes
from navigate.core.node_report import NodeReport
from navigate.core.nodes.report import Report
from navigate.core.nodes.vessel import Vessel
from navigate.output import report_writer
from navigate.output.report_writer import ROW_RESULT, _prepare_export, write_report


def _node_report(attribute="Lifetime"):
    report = NodeReport()
    report.add_property(attribute, ReportReduceID.NONE)
    return report


class TestPrepareExport:
    def test_matched_name_exports(self):
        nodes = {"vessel": Vessel("vessel")}

        export = _prepare_export(nodes, {"vessel": _node_report()}, "output", "Vessels")

        assert set(export) == {"vessel"}
        assert export["vessel"][0] == ["Lifetime"]

    def test_unmatched_name_warns_and_skips(self, caplog):
        nodes = {"vessel": Vessel("vessel")}

        with caplog.at_level(logging.WARNING):
            export = _prepare_export(
                nodes, {"ghost": _node_report()}, "output", "Vessels"
            )

        assert export == {}
        assert (
            "Report 'output': property request 'ghost' on sheet 'Vessels'"
            in caplog.text
        )


class TestWriteReportErrorContainment:
    @staticmethod
    def _report_and_manager():
        report = Report("output")
        report.add_fleet_property("fleet", "Lifetime")
        report.add_vessel_property("vessel", "Lifetime")

        manager = SimpleNamespace(
            name="global", nodes=Nodes(fleets={"fleet": None}, vessels={"vessel": None})
        )
        return report, manager

    def test_failed_xlsx_sheet_is_skipped_and_the_report_still_saves(
        self, tmp_path, caplog, monkeypatch
    ):
        report, manager = self._report_and_manager()

        def fail_on_fleets(ws, nodes, extraction_dict, report_name):
            if ws.title == "Fleets":
                raise RuntimeError("boom")
            ws.cell(row=ROW_RESULT, column=3).value = "exported"

        monkeypatch.setattr(report_writer, "export_properties_xlsx", fail_on_fleets)
        dateline = np.array(["2030-01-01", "2031-01-01"], dtype="datetime64[D]")

        with caplog.at_level(logging.WARNING):
            write_report(report, manager, str(tmp_path), "deck", dateline)

        wb = xl.load_workbook(tmp_path / "deck_output.xlsx")
        assert wb["Vessels"].cell(row=ROW_RESULT, column=3).value == "exported"
        assert "Report 'output': Failed to export 'fleets': boom" in caplog.text
        assert "Report 'output': Completed with 1 sheet error(s)." in caplog.text

    def test_failed_csv_sheet_is_skipped_and_the_report_still_saves(
        self, tmp_path, caplog, monkeypatch
    ):
        report, manager = self._report_and_manager()
        report.set_file_format("CSV")

        def fail_on_fleets(sheet_name, nodes, extraction_dict, report_name, csv_data):
            if sheet_name == "Fleets":
                raise RuntimeError("boom")
            csv_data[sheet_name] = {
                "headers": ["marker"],
                "columns": [np.array([1.0, 2.0])],
            }

        monkeypatch.setattr(report_writer, "export_properties_csv", fail_on_fleets)
        dateline = np.array(["2030-01-01", "2031-01-01"], dtype="datetime64[D]")

        with caplog.at_level(logging.WARNING):
            write_report(report, manager, str(tmp_path), "deck", dateline)

        assert (tmp_path / "deck_output_Vessels.csv").exists()
        assert not (tmp_path / "deck_output_Fleets.csv").exists()
        assert "Report 'output': Failed to export 'fleets': boom" in caplog.text
        assert "Report 'output': Completed with 1 sheet error(s)." in caplog.text
